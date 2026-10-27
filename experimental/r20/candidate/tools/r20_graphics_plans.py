"""R2.0: independent graphical storyboard from a SINGLE editorial DOCX.
No montage, source clips, timecode, ASR, or Premiere dependency. No cloud I/O.
All timing is SCRIPT_PREVIEW / EDITORIAL_PREVIEW, never speech-sync.
"""
from __future__ import annotations
import argparse, hashlib, json, math, re
from pathlib import Path
from docx import Document
from reportaj_r1_parse_docx import extract
from r193_docx_reader import read as read_r193
from r161_prepare_nametags import card_for
from r171_layout import split_cards, highlight_by_words, clean
from r153_prepare import split_balanced

SCHEMA='aive-r20-graphics-v1'
SEPARATOR=.75

def duration_for(text, minimum=1.15, maximum=8.0, cps=14.0):
    return round(min(maximum, max(minimum, .6+len(text)/cps)), 3)

def _extract(docx):
    first=extract(docx)
    second=read_r193(docx)
    a=first['segments']; b=second['segments']
    if len(b)>len(a)>0:
        raise ValueError('Cititorul nou găsește %d sincronuri, dar cel legacy doar %d: posibil text lipsă.'%(len(b),len(a)))
    if a:
        # R1 parser can accidentally append an isolated OUT timecode as spoken text.
        # Remove ONLY a standalone timecode paragraph, never inline editorial text.
        paragraph_text=[p.text.strip() for p in Document(docx).paragraphs]
        for j,seg in enumerate(a):
            start=seg.get('paragraph_start',-1)
            end=a[j+1].get('paragraph_start',len(paragraph_text)) if j+1<len(a) else len(paragraph_text)
            for para in paragraph_text[max(start,0):end]:
                if re.fullmatch(r'\d{1,2}[.:]\d{2}',para) and seg['text'].endswith(' '+para):
                    seg['text']=seg['text'][:-(len(para)+1)].rstrip()
        segments=a
    elif b:
        segments=[{'ordinal':i+1, 'paragraph_start':x['paragraph'], 'paragraph_end':x['paragraph'],
                   'source_code':x['source_code'], 'speaker':x['speaker'], 'role':x['role'],
                   'text':x['text'], 'reporter_italic':x.get('reporter_italic',False)} for i,x in enumerate(b)]
    else:raise ValueError('Nu am identificat declarații în DOCX.')
    captions=first['graphics'] or second['captions']
    if len(second['captions'])>len(first['graphics'])>0:
        raise ValueError('Cititorul nou vede mai multe cartoane decât cel legacy; revizuire obligatorie.')
    return segments,captions,first


def _editorial_markers(docx:Path):
    """Collect identity-boundary markers without requiring standalone BOLD 4-digit codes.

    Reportage scripts may place SINCRON/AMB, name, role, source and TC on
    separate *lines of the same paragraph*. The old R1.6 finder only sees an
    entirely bold, four-digit *paragraph*, which misses these layouts.
    """
    markers=[]
    for index, paragraph in enumerate(Document(docx).paragraphs):
        for line in paragraph.text.splitlines():
            line=line.strip()
            if re.match(r'^(?:SINCRON|AMB)\b\s*[:–—-]',line,re.I):
                markers.append((index,'editorial',None))
            elif re.fullmatch(r'\d{4}',line):
                markers.append((index,'source',line))
    return markers


def _name_blocks(segments,markers):
    """Yield (segment, card, block number) for consecutive editorial blocks.

    Always count a reintroduced speaker as a new block after an intervening
    speaker or a fresh heading. Never globally deduplicate by speaker name.
    If the source has no explicit heading, use only *consecutive* identities
    present in the already approved DOCX parser, not fabricated headings.
    """
    prior_identity=None;prior_marker=None;block_number=0
    for seg in segments:
        card=card_for(seg)
        if card is None:
            # An unidentified segment interrupts contiguous identity grouping.
            prior_identity=None;prior_marker=None
            yield seg,None,None
            continue
        index=int(seg['paragraph_start'])
        code=seg.get('source_code')
        before=[m for m in markers if m[0]<=index]
        if before:
            # Preserve source-code checks: a contradicting explicit SOURCE
            # marker may not be silently mapped onto a named speaker.
            last_source=next((m for m in reversed(before) if m[1]=='source'),None)
            if last_source and code and last_source[2]!=code:
                raise ValueError('NameTag REVIEW: cod sursă DOCX %s ≠ %s la CUT-%02d' %
                                 (last_source[2],code,seg['ordinal']))
            marker=(before[-1][0],before[-1][1])
        else:
            marker=None
        identity=(code,card[0],card[1],bool(card[2]))
        if identity!=prior_identity or marker!=prior_marker:
            block_number+=1
        prior_identity=identity;prior_marker=marker
        yield seg,card,block_number


def build(docx:Path, module:str, separator:float=SEPARATOR):
    docx=Path(docx).resolve()
    if not docx.is_file() or docx.suffix.lower()!='.docx':raise ValueError('DOCX invalid / fișier lipsă')
    if module not in {'02','03','04'}:raise ValueError('Modul neacceptat (02/03/04)')
    if not math.isfinite(separator) or not 0<=separator<=2:raise ValueError('Separator invalid')
    segments,captions,source=_extract(docx)
    items=[]; skipped=[]; cursor=0.0; count=0
    def add(group,text,seconds,other=None):
        nonlocal cursor,count
        if not text.strip():raise ValueError('Element grafic gol')
        group=str(group)
        if items and items[-1]['groupId']!=group:cursor+=separator
        count+=1
        x={'id':'R20-%s-%04d'%(module,count),'groupId':group,'text':text,
           'startSeconds':round(cursor,3),'durationSeconds':round(seconds,3)}
        if other:x.update(other)
        items.append(x)
        cursor=round(cursor+seconds,3)
    if module=='02':
        for seg in segments:
            original=clean(seg['text'])
            chunks=split_balanced(original,74)
            if not chunks or clean(' '.join(chunks))!=original:
                raise ValueError('Text de subtitrare pierdut la CUT-%02d'%seg['ordinal'])
            group='CUT-%02d'%seg['ordinal']
            style='Italic' if seg.get('reporter_italic') is True else 'Regular'
            for chunk in chunks:
                add(group,chunk,duration_for(chunk,minimum=.9,maximum=6.0,cps=13.0),
                    {'fontStyle':style,'sourceParagraph':seg['paragraph_start']})
    elif module=='03':
        markers=_editorial_markers(docx)
        recorded=set()
        for seg,name,block in _name_blocks(segments,markers):
            if name is None:
                skipped.append({'paragraph':seg['paragraph_start'],
                                'cutId':'CUT-%02d'%seg['ordinal'],
                                'reason':'Lipsesc nume/funcție din DOCX'})
                continue
            if block in recorded:continue
            recorded.add(block)
            person,title,anonymous=name
            add('BLOCK-'+str(block),person,duration_for(person+' '+title,minimum=3.3,maximum=5.5,cps=16),
                {'name':person,'title':title,'anonymous':bool(anonymous),'sourceParagraph':seg['paragraph_start']})
    else:
        if not captions:raise ValueError('Nu există cartoane editoriale în DOCX')
        for j,block in enumerate(captions):
            original=clean(block['text'])
            pages=split_cards(original)
            if not pages or clean(' '.join(raw for raw,_ in pages))!=original:
                raise ValueError('Textul cartonului a fost modificat: bloc '+str(j+1))
            for k,(raw,slots) in enumerate(pages):
                bold,_=highlight_by_words(raw,slots)
                add('VO-'+str(j+1),raw,duration_for(raw,minimum=3,maximum=9,cps=13.5),
                    {'physicalLines':slots,'textLine0':slots[0],'textLines123':'\r'.join(slots[1:]),
                     'boldWords':bold,'colorScheme':'Orange','sourceParagraph':block['paragraph'],
                     'page':k+1,'pageCount':len(pages)})
    if not items:raise ValueError('Nu există elemente valide pentru categoria '+module)
    if module=='03' and skipped: # disallow silent omitted named blocks
        raise ValueError('NameTag incomplet: %d identificări fără nume/funcție. Revizuire editorială obligatorie.'%len(skipped))
    return {'schema':SCHEMA,'module':module,'mode':'independent',
            'timing':'SCRIPT_PREVIEW' if module=='02' else 'EDITORIAL_PREVIEW',
            'docx_sha256':hashlib.sha256(docx.read_bytes()).hexdigest(),
            'sourceDocx':docx.name,'separatorSeconds':separator,'items':items,
            'totalItems':len(items),'groups':len({a['groupId'] for a in items}),
            'skipped':skipped,'containsAudioSync':False,'generatorVersion':'R2.0-candidate'}

def main():
    p=argparse.ArgumentParser()
    p.add_argument('--docx',required=True,type=Path);p.add_argument('--module',required=True,choices=('02','03','04'))
    p.add_argument('--output',required=True,type=Path);p.add_argument('--separator',default=.75,type=float)
    x=p.parse_args()
    plan=build(x.docx,x.module,x.separator)
    x.output.parent.mkdir(parents=True,exist_ok=True)
    tmp=x.output.with_suffix('.json.tmp')
    tmp.write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')
    tmp.replace(x.output)
    print(json.dumps({'ok':True,'plan':str(x.output),'count':plan['totalItems'],'groups':plan['groups'],
                      'docx_sha256':plan['docx_sha256'],'timing':plan['timing']},ensure_ascii=False))
if __name__=='__main__':main()
