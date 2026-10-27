"""R2.0.3: incomplete editorial identities require EXPLICIT skip approval in CEP.

Synthetic input only. No user DOCX or private transcript enters public GitHub.
"""
import contextlib, io, json, sys, tempfile, unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from test_r201_nametag import planner, seg

class TestR203ReviewGate(unittest.TestCase):
    def setUp(self):
        self.t=tempfile.TemporaryDirectory()
        self.doc=Path(self.t.name)/'synthetic.docx'
        d=Document()
        for i in range(10):d.add_paragraph('Fragment editorial de verificare '+str(i))
        d.save(self.doc)
    def tearDown(self):self.t.cleanup()
    def segments(self):
        return [seg(i+1,i,'1234','Persoana '+str(i+1) if i not in (4,8) else None,
                    'Funcție '+str(i+1) if i not in (4,8) else None)
                for i in range(10)]
    def test_eight_valid_two_review_without_fake_names(self):
        with patch.object(planner,'_extract',return_value=(self.segments(),[],{})):
            result=planner.build(self.doc,'03')
        self.assertEqual(result['totalItems'],8)
        self.assertEqual(result['groups'],8)
        self.assertEqual([x['cutId'] for x in result['skipped']],['CUT-05','CUT-09'])
        self.assertEqual([x['paragraph'] for x in result['skipped']],[4,8])
        self.assertTrue(result['reviewRequired'])
        self.assertEqual(result['reviewPolicy'],'EXPLICIT_SKIP_APPROVAL')
        self.assertEqual(result['timing'],'EDITORIAL_PREVIEW')
        self.assertTrue(all(x['name'] not in ('',None) and x['title'] not in ('',None) for x in result['items']))
        self.assertEqual([x['kind'] for x in result['skipped']],['REVIEW_MISSING_IDENTITY']*2)
        self.assertFalse(any('CUT-05' in x['groupId'] or 'CUT-09' in x['groupId'] for x in result['items']))
    def test_explicit_ambience_9974_auto_skip_with_eight_names(self):
        doc=Document()
        segments=[]
        for i in range(10):
            ambient=i in (4,8)
            code='9974' if ambient else str(9900+i)
            para=('AMB: Vânzare plăcinte împreună cu o colegă' if ambient else
                  'SINCRON: Persoana %s, Elev' % (i+1))
            doc.add_paragraph(para+'\\n'+code+'\\n00.08–00.16')
            segments.append(seg(i+1,i,code,
                None if ambient else 'Persoana '+str(i+1),
                None if ambient else 'Elev'))
        path=Path(self.t.name)/'realistic_amb.docx';doc.save(path)
        with patch.object(planner,'_extract',return_value=(segments,[],{})):
            p=planner.build(path,'03')
        self.assertEqual(p['totalItems'],8)
        self.assertEqual([x['cutId'] for x in p['skipped']],['CUT-05','CUT-09'])
        self.assertEqual([x['sourceCode'] for x in p['skipped']],['9974','9974'])
        self.assertEqual([x['kind'] for x in p['skipped']],['SKIP_AMBIENCE','SKIP_AMBIENCE'])
        self.assertEqual(p['reviewPolicy'],'AUTO_AMBIENCE_SKIP')
        self.assertFalse(p['reviewRequired'])
        panel=(Path(__file__).resolve().parents[1]/'cep/r20_panel.js').read_text('utf8')
        self.assertIn("if(plan.reviewRequired)",panel)
        self.assertIn("SKIP_AMBIENCE",panel)
    def test_no_auto_skip_if_heading_missing_or_source_conflicts(self):
        with self.assertRaisesRegex(ValueError,'REVIEW: cod sursă'):
            planner._skip_identity_reason(seg(5,4,'9981',None,None),[(3,'ambience',None),(3,'source','9974')])
        kind,_=planner._skip_identity_reason(seg(5,4,'9974',None,None),[])
        self.assertEqual(kind,'REVIEW_MISSING_IDENTITY')
    def test_cli_persists_review_identically(self):
        target=Path(self.t.name)/'plan.json'
        with patch.object(planner,'_extract',return_value=(self.segments(),[],{})):
            with patch.object(sys,'argv',['r20_graphics_plans.py','--docx',str(self.doc),'--module','03','--output',str(target)]):
                log=io.StringIO()
                with contextlib.redirect_stdout(log):planner.main()
        signal=json.loads(log.getvalue())
        plan=json.loads(target.read_text('utf8'))
        self.assertTrue(signal['ok'])
        self.assertEqual(signal['skipped'],plan['skipped'])
        self.assertTrue(signal['reviewRequired'])
        self.assertEqual(signal['count'],8)
        self.assertEqual(signal['docx_sha256'],plan['docx_sha256'])
    def test_no_missing_no_extra_confirmation(self):
        all_named=[seg(1,0,'1234','Ana','Director'),seg(2,1,'5678','Ion','Profesor')]
        with patch.object(planner,'_extract',return_value=(all_named,[],{})):
            p=planner.build(self.doc,'03')
        self.assertEqual(p['skipped'],[])
        self.assertFalse(p['reviewRequired'])
        self.assertEqual(p['reviewPolicy'],'NONE')
    def test_all_unnamed_still_fails_closed(self):
        all_blank=[seg(1,0,'1234',None,None),seg(2,1,'5678',None,None)]
        with patch.object(planner,'_extract',return_value=(all_blank,[],{})):
            with self.assertRaisesRegex(ValueError,'Nu există elemente valide'):
                planner.build(self.doc,'03')
    def test_synced_executor_and_legacy_untouched(self):
        p=Path(__file__).resolve().parents[1]
        panel=(p/'cep/r20_panel.js').read_text('utf8')
        self.assertIn("if(!window.confirm(question))",panel)
        self.assertIn("Nicio inserare V4",panel)
        self.assertIn("REVIEW APROBAT",panel)
        self.assertLess(panel.index("if(!window.confirm(question))"),panel.index("evalR20('validate'"))
        self.assertIn("if(mode.value==='independent')run(k);else if(old)old.call(btn);",panel)
if __name__=='__main__':unittest.main(verbosity=2)
