"""R2.0.1 isolated NameTag parser regressions; no private legacy code in PUBLIC repo.
Import stubs expose only the narrow stable dependencies used by the changed helper.
The real 12-test suite was also executed offline against the installed source copy.
"""
import sys,types,tempfile,unittest
from pathlib import Path
from docx import Document

def fake(name,**attrs):
    m=types.ModuleType(name)
    for k,v in attrs.items():setattr(m,k,v)
    sys.modules[name]=m

def card_for(seg):
    speaker=(seg.get('speaker') or '').strip()
    role=(seg.get('role') or '').strip()
    return (speaker,role,False) if speaker and role else None

fake('reportaj_r1_parse_docx',extract=lambda p:None)
fake('r193_docx_reader',read=lambda p:None)
fake('r161_prepare_nametags',card_for=card_for)
fake('r171_layout',split_cards=lambda t:[],highlight_by_words=lambda a,b:({},''),clean=lambda t:t)
fake('r153_prepare',split_balanced=lambda s,n:[s])
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'tools'))
import r20_graphics_plans as planner

def seg(i,p,code,name,role,text='Text editorial'):
    return {'ordinal':i,'paragraph_start':p,'source_code':code,
            'speaker':name,'role':role,'text':text}

class TestR201(unittest.TestCase):
    def setUp(self):self.t=tempfile.TemporaryDirectory()
    def tearDown(self):self.t.cleanup()
    def test_inline_merged_headers(self):
        doc=Document();doc.add_paragraph('SCRIPT')
        doc.add_paragraph('SINCRON: Ana Ștefănescu, Directoare\n1234\n0:01-0:08\nPrima declarație despre școală.')
        doc.add_paragraph('0:09-0:13\nA doua declarație a aceleiași persoane.')
        doc.add_paragraph('AMB: Ion Popescu, Profesor\n5678\n0:02-0:05\nÎn atelier elevii lucrează.')
        doc.add_paragraph('SINCRON: Ana Ștefănescu, Directoare\n1234\n0:16-0:20\nAna a revenit.')
        p=Path(self.t.name)/'merged.docx';doc.save(p)
        markers=planner._editorial_markers(p)
        self.assertEqual(markers,[(1,'editorial',None),(1,'source','1234'),
                                  (3,'ambience',None),(3,'source','5678'),
                                  (4,'editorial',None),(4,'source','1234')])
        segments=[seg(1,1,'1234','Ana Ștefănescu','Directoare'),
                  seg(2,2,'1234','Ana Ștefănescu','Directoare'),
                  seg(3,3,'5678','Ion Popescu','Profesor'),
                  seg(4,4,'1234','Ana Ștefănescu','Directoare')]
        old=planner._extract
        try:
            planner._extract=lambda doc: (segments,[],{})
            plan=planner.build(p,'03')
        finally:
            planner._extract=old
        self.assertEqual([x['name'] for x in plan['items']],
                         ['Ana Ștefănescu','Ion Popescu','Ana Ștefănescu'])
        self.assertEqual(plan['totalItems'],3)
        self.assertEqual(plan['timing'],'EDITORIAL_PREVIEW')
    def test_missing_first_header(self):
        src=[seg(1,5,'9981','Elena','Director'),seg(2,7,'9981','Elena','Director'),
             seg(3,10,'9974','Ion','Profesor'),seg(4,12,'9981','Elena','Director')]
        self.assertEqual([r[2] for r in planner._name_blocks(src,[])],[1,1,2,3])
    def test_repeated_identity_new_header(self):
        src=[seg(i+1,p,'1234','Ana','Director') for i,p in enumerate([3,4,7])]
        self.assertEqual([r[2] for r in planner._name_blocks(
            src,[(2,'editorial',None),(6,'editorial',None)])],[1,1,2])
    def test_conflicting_source_rejected(self):
        src=[seg(1,5,'9981','Ana','Director')]
        with self.assertRaisesRegex(ValueError,'REVIEW: cod sursă'):
            list(planner._name_blocks(src,[(4,'source','9974')]))
    def test_unknown_identity_no_fabrication(self):
        src=[seg(1,5,'1234','Ana','Director'),seg(2,6,'1234',None,None),
             seg(3,7,'1234','Ana','Director')]
        self.assertEqual([r[2] for r in planner._name_blocks(src,[])],[1,None,2])
if __name__=='__main__':
    unittest.main(verbosity=2)
