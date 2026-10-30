"""R236 private-script-structure regression with SYNTHETIC text only:
six source/name/role blocks, 0007 vs 0034, Medic role, mixed strike.
No real personal DOCX text is checked into GitHub.
"""
import sys,tempfile,types
from pathlib import Path
from docx import Document
ROOT=Path('/tmp/r236_runtime/tools');sys.path.insert(0,str(ROOT))
def clean(x):return ' '.join(x.split())
def pages(x):return []
def split(txt,limit=74):
 result=[];part=''
 for w in txt.split():
  if part and len(part)+len(w)+1>limit:result.append(part);part=w
  else:part=(part+' '+w).strip()
 if part:result.append(part)
 return result
sys.modules["reportaj_r1_parse_docx"]=types.SimpleNamespace(extract=lambda x:{"segments":[],"graphics":[]})
sys.modules["r193_docx_reader"]=types.SimpleNamespace(read=lambda x:{"segments":[],"captions":[]})
sys.modules["r161_prepare_nametags"]=types.SimpleNamespace(card_for=lambda x:None)
sys.modules["r171_layout"]=types.SimpleNamespace(clean=clean)
sys.modules["r204_carton_layout"]=types.SimpleNamespace(build_pages=pages)
sys.modules["r207_docx_strike"]=types.SimpleNamespace(extract_strike_policy=lambda *args:None)
sys.modules["r153_prepare"]=types.SimpleNamespace(split_balanced=lambda x,limit=74:split(x,limit))
import r20_graphics_plans as G
from r236_name_guard import check
from r230_universal import build as universal_build
with tempfile.TemporaryDirectory() as f:
 p=Path(f)/"synthetic Balatina editorial.docx"
 d=Document()
 d.add_paragraph("Medicii din comunitate continuă să lucreze.")
 d.add_paragraph("Ambianță")
 d.add_paragraph("0012")
 d.add_paragraph("00.09 Aceasta este o replică înregistrată ambiental. 00.20")
 codes=["0007","0007","0034","0034","0007","0007"]
 who=["Nina Exemplu","Nina Exemplu","Eugenia Exemplu","Eugenia Exemplu","Nina Exemplu","Nina Exemplu"]
 roles=["Șefa Centrului de Sănătate","Șefa Centrului de Sănătate",
        "Medic de familie","Medic de familie","Șefa Centrului de Sănătate","Șefa Centrului de Sănătate"]
 for i,(c,n,r) in enumerate(zip(codes,who,roles)):
  d.add_paragraph("MEDICII DIN SAT LUCREAZĂ PENTRU PACIENȚII LOR")
  d.add_paragraph(c);d.add_paragraph(n);d.add_paragraph(r)
  para=d.add_paragraph()
  para.add_run("1.32 Declarația principală a persoanei și răspunsul ei. ")
  strike=para.add_run("Fragmentul acesta nu trebuie montat.")
  if i==2:strike.font.strike=True
  para.add_run(" 1.51")
 d.save(p)
 pure=universal_build(p,'03')
 assert pure['totalItems']==6,[(x.get('name'),x.get('title')) for x in pure['items']]
 assert [x['name'] for x in pure['items']]==who
 assert all(x['title']==r for x,r in zip(pure['items'],roles))
 guard=check(p,pure)
 assert guard['verified']==6 and guard['explicitAnchors']==6
 old=G._r230_legacy_build
 try:
  def legacy(*args,**kwargs):
   raise ValueError('NameTag REVIEW: cod sursă DOCX 0034 ≠ 0007 la CUT-03')
  G._r230_legacy_build=legacy
  result=G.build(p,'03')
  assert result['totalItems']==6 and result['sourceNameReview']['verified']==6
 finally:G._r230_legacy_build=old
 print('PASS six explicit NameTags recovered despite legacy CODE mismatch; no identities fabricated')
 assert universal_build(p,'02',omit_struck=True)['strikeCutsChanged']>=1
 print('PASS strike handling remains on V3; Medic de familie explicit role preserved')
