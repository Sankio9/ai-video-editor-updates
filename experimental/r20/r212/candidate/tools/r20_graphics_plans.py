"""R2.0: independent graphical storyboard from a SINGLE editorial DOCX.
No montage, source clips, timecode, ASR, or Premiere dependency. No cloud I/O.
All timing is SCRIPT_PREVIEW / EDITORIAL_PREVIEW, never speech-sync.
"""
from __future__ import annotations
import argparse, hashlib, json, math, re
from pathlib import Path
from docx import Document
from reportaj_r1_parse_docx import extract
from r193_docx_reader import read as read_r193
from r161_prepare_nametags import card_for
from r171_layout import clean
from r204_carton_layout import build_pages
from r207_docx_strike import extract_strike_policy
from r153_prepare import split_balanced

SCHEMA='aive-r20-graphics-v1'
SEPARATOR=.75

def duration_for(text, minimum=1.15, maximum=8.0, cps=14.0):
    return round(min(maximum, max(minimum, .6+len(text)/cps)), 3)


_SENTENCE_END=re.compile(r'[.!?…][”»"’]*$')
# Function words are forbidden at the END of a subtitle when more text follows.
# This also protects "pot să", "nu stea" and "să informăm".
_DANGLING=frozenset((
    'să','și','dar','iar','sau','că','de','la','cu','în','un','o','pe',
    'prin','pentru','din','care','al','a','ai','ale','nu','mai','pot'
))
_CLAUSE_START=frozenset(('dar','iar','însă','totuși','deci','deoarece','fiindcă'))

def split_subtitles_readable(text, limit=74):
    """Balanced, punctuation-aware captions; no rewriting or word splitting.

    Minimize captions first, but allow ONE extra cue if the minimum would
    force a stranded sentence starter or an isolated very short remainder.
    All output keeps exact word order, punctuation and original wording.
    """
    text=clean(text)
    if not text:return []
    if not isinstance(limit,int) or limit<20:
        raise ValueError('Limită subtitrare invalidă')
    words=text.split(' ')
    n=len(words)
    if any(len(word)>limit for word in words):
        raise ValueError('Un cuvânt depășește limita de subtitrare')
    if len(text)<=limit:return [text]

    prefix=[0]
    for word in words:prefix.append(prefix[-1]+len(word))
    def charlen(i,j):
        return prefix[j]-prefix[i]+j-i-1

    # Minimum number of non-split whole-word cues: unambiguous lower bound.
    minimum=0;pos=0
    while pos<n:
        right=pos+1
        while right<n and charlen(pos,right+1)<=limit:right+=1
        pos=right;minimum+=1

    inf=float('inf')
    choices=[]
    for parts in range(minimum,min(n,minimum+1)+1):
        target=float(len(text))/parts
        cost=[[inf]*(n+1) for _ in range(parts+1)]
        parent=[[None]*(n+1) for _ in range(parts+1)]
        cost[0][0]=0.0

        def edge(i,j):
            count=charlen(i,j)
            penalty=110.0*((count-target)/max(20.0,target))**2
            minimum_readable=min(31.0,target*.75)
            if count<minimum_readable:
                penalty+=120.0*((minimum_readable-count)/max(1.0,minimum_readable))**2
            # A cue should not merge the tail of an entire sentence
            # with a substantial fragment of the next sentence. Keep
            # semantic boundaries whenever balancing can afford it.
            for q in range(i+1,j):
                if (_SENTENCE_END.search(words[q-1]) and
                    charlen(i,q)>=18 and charlen(q,j)>=18):
                    penalty+=75.0
            if j==n:return penalty

            ending=words[j-1].strip('.,!?;:„”«»').casefold()
            if ending in _DANGLING:
                return inf  # never end a cue with "Să", "dar", "pot", etc.

            previous=words[j-1];following=words[j].casefold().strip('„«')
            if _SENTENCE_END.search(previous):
                penalty-=40.0
            elif previous.endswith((',', ';', ':')):
                penalty-=14.0
            elif following in _CLAUSE_START:
                penalty-=8.0
            else:
                penalty+=6.0

            # No new sentence may begin with only 1-2 words at the end of
            # the preceding cue. Prefer moving all of it into the next cue.
            for q in range(max(i+1,j-2),j):
                if _SENTENCE_END.search(words[q-1]) and j-q<=2:
                    return inf
            if following in ('mai','departe'):penalty+=12.0
            return penalty

        for section in range(1,parts+1):
            for j in range(section,n+1):
                for i in range(section-1,j):
                    if cost[section-1][i]==inf:continue
                    if charlen(i,j)>limit:continue
                    value=cost[section-1][i]+edge(i,j)
                    if value<cost[section][j]:
                        cost[section][j]=value
                        parent[section][j]=i
        if parent[parts][n] is None:continue
        cues=[];j=n
        for section in range(parts,0,-1):
            i=parent[section][j]
            cues.append(' '.join(words[i:j]))
            j=i
        cues.reverse()
        # 24 units discourages unnecessary captions, except when a short
        # orphan would otherwise be forced by a full 74-char first cue.
        choices.append((cost[parts][n]+24.0*(parts-minimum),parts,cues))

    if not choices:
        # The constraints should be satisfiable for ordinary prose. Never
        # silently generate a caption with orphaned function words.
        raise ValueError('Textul nu poate fi împărțit lizibil fără resturi izolate')
    result=min(choices,key=lambda option:(option[0],option[1]))[2]
    if any(not cue or len(cue)>limit for cue in result):
        raise ValueError('Subtitrare depășește limita sau este goală')
    if clean(' '.join(result))!=text:
        raise ValueError('Subtitrări: pierdere text editorial')
    return result


def _extract(docx):
    first=extract(docx)
    second=read_r193(docx)
    a=first['segments']; b=second['segments']
    if len(b)>len(a)>0:
        raise ValueError('Cititorul nou găsește %d sincronuri, dar cel legacy doar %d: posibil text lipsă.'%(len(b),len(a)))
    if a:
        # R1 parser can accidentally append an isolated OUT timecode as spoken text.
        # Remove ONLY a standalone timecode paragraph, never inline editorial text.
        paragraph_text=[p.text.strip() for p in Document(docx).paragraphs]
        for j,seg in enumerate(a):
            start=seg.get('paragraph_start',-1)
            end=a[j+1].get('paragraph_start',len(paragraph_text)) if j+1<len(a) else len(paragraph_text)
            for para in paragraph_text[max(start,0):end]:
                if re.fullmatch(r'\d{1,2}[.:]\d{2}',para) and seg['text'].endswith(' '+para):
                    seg['text']=seg['text'][:-(len(para)+1)].rstrip()
        segments=a
    elif b:
        segments=[{'ordinal':i+1, 'paragraph_start':x['paragraph'], 'paragraph_end':x['paragraph'],
                   'source_code':x['source_code'], 'speaker':x['speaker'], 'role':x['role'],
                   'text':x['text'], 'reporter_italic':x.get('reporter_italic',False)} for i,x in enumerate(b)]
    else:raise ValueError('Nu am identificat declarații în DOCX.')
    captions=first['graphics'] or second['captions']
    if len(second['captions'])>len(first['graphics'])>0:
        raise ValueError('Cititorul nou vede mai multe cartoane decât cel legacy; revizuire obligatorie.')
    return segments,captions,first


def _editorial_markers(docx:Path):
    """Collect identity-boundary markers without requiring standalone BOLD 4-digit codes.

    Reportage scripts may place SINCRON/AMB, name, role, source and TC on
    separate *lines of the same paragraph*. The old R1.6 finder only sees an
    entirely bold, four-digit *paragraph*, which misses these layouts.
    """
    markers=[]
    for index, paragraph in enumerate(Document(docx).paragraphs):
        for line in paragraph.text.splitlines():
            line=line.strip()
            heading=re.match(r'^(SINCRON|AMB)\b\s*[:–—-]',line,re.I)
            if heading:
                markers.append((index,'ambience' if heading.group(1).upper()=='AMB' else 'editorial',None))
            elif re.fullmatch(r'\d{4}',line):
                markers.append((index,'source',line))
    return markers


def _skip_identity_reason(seg, markers):
    """Editorially safe: an explicit AMB source is NOT a speaker NameTag.

    Do not silently treat unnamed SINCRON as ambience. Unnamed interview
    remains REVIEW with an explicit user decision. Check source-code agreement
    for both named and anonymous segments.
    """
    start=int(seg['paragraph_start'])
    relevant=[m for m in markers if m[0]<=start]
    last_heading=next((m for m in reversed(relevant) if m[1] in ('ambience','editorial')),None)
    last_source=next((m for m in reversed(relevant) if m[1]=='source'),None)
    code=seg.get('source_code')
    if last_source and code and last_source[2]!=code:
        raise ValueError('NameTag REVIEW: cod sursă DOCX %s ≠ %s la CUT-%02d' %
                         (last_source[2],code,seg['ordinal']))
    if (last_heading and last_heading[1]=='ambience' and
            last_source and last_source[0]>=last_heading[0] and
            (not code or last_source[2]==code)):
        return ('SKIP_AMBIENCE','AMB editorial fără identificare NameTag; nu este burtieră')
    return ('REVIEW_MISSING_IDENTITY','SINCRON/identificare fără nume și funcție verificabile')

def _name_blocks(segments,markers):
    """Yield (segment, card, block number) for consecutive editorial blocks.

    Always count a reintroduced speaker as a new block after an intervening
    speaker or a fresh heading. Never globally deduplicate by speaker name.
    If the source has no explicit heading, use only *consecutive* identities
    present in the already approved DOCX parser, not fabricated headings.
    """
    prior_identity=None;prior_marker=None;block_number=0
    for seg in segments:
        card=card_for(seg)
        if card is None:
            # An unidentified segment interrupts contiguous identity grouping.
            prior_identity=None;prior_marker=None
            yield seg,None,None
            continue
        index=int(seg['paragraph_start'])
        code=seg.get('source_code')
        before=[m for m in markers if m[0]<=index]
        if before:
            # Preserve source-code checks: a contradicting explicit SOURCE
            # marker may not be silently mapped onto a named speaker.
            last_source=next((m for m in reversed(before) if m[1]=='source'),None)
            if last_source and code and last_source[2]!=code:
                raise ValueError('NameTag REVIEW: cod sursă DOCX %s ≠ %s la CUT-%02d' %
                                 (last_source[2],code,seg['ordinal']))
            marker=(before[-1][0],before[-1][1])
        else:
            marker=None
        identity=(code,card[0],card[1],bool(card[2]))
        if identity!=prior_identity or marker!=prior_marker:
            block_number+=1
        prior_identity=identity;prior_marker=marker
        yield seg,card,block_number

def build(docx:Path, module:str, separator:float=SEPARATOR, omit_struck:bool=False):
    docx=Path(docx).resolve()
    if not docx.is_file() or docx.suffix.lower()!='.docx':raise ValueError('DOCX invalid / fișier lipsă')
    if module not in {'02','03','04'}:raise ValueError('Modul neacceptat (02/03/04)')
    if not math.isfinite(separator) or not 0<=separator<=2:raise ValueError('Separator invalid')
    segments,captions,source=_extract(docx)
    if type(omit_struck) is not bool:raise ValueError('Politica DOCX Strikethrough trebuie booleană')
    strike=extract_strike_policy(docx,segments,omit_struck) if module=='02' else None
    # A different/old Python detector may incorrectly return 0 removals.
    # The PLAN, not just the helper, must fail before any Premiere command.
    if module=='02' and omit_struck and (strike['policy']!='OMIT_STRUCK' or
            strike['removedCharacters']<=0 or strike['cutsChanged']<=0):
        raise ValueError('STRIKETHROUGH_NOT_DETECTED: verificarea DOCX nu a eliminat niciun caracter. '
                         'OPRIT ÎNAINTE de inserarea pe V3.')
    strike_audit=[]
    if module=='02':
        if len(strike['texts'])!=len(segments):
            raise ValueError('STRIKE_ANALYSIS_INCOMPLETE: lipsesc CUT-uri din analiza Word')
        for seg,filtered in zip(segments,strike['texts']):
            before=' '.join(str(seg['text']).split())
            after=' '.join(str(filtered).split())
            if before!=after:
                if not omit_struck:raise ValueError('STRIKE_POLICY_CONTRADICTORY')
                strike_audit.append({'cutId':'CUT-%02d'%seg['ordinal'],
                                     'paragraph':seg['paragraph_start'],
                                     'before':before,'after':after})
        if omit_struck and not strike_audit:
            raise ValueError('STRIKE_ANALYSIS_INCOMPLETE: bifa ON, niciun text schimbat')
        if omit_struck and len(strike_audit)!=strike['cutsChanged']:
            raise ValueError('STRIKE_ANALYSIS_INCOMPLETE: numărul CUT-urilor diferă')
    items=[]; skipped=[]; cursor=0.0; count=0
    def add(group,text,seconds,other=None):
        nonlocal cursor,count
        if not text.strip():raise ValueError('Element grafic gol')
        group=str(group)
        if items and items[-1]['groupId']!=group:cursor+=separator
        count+=1
        x={'id':'R20-%s-%04d'%(module,count),'groupId':group,'text':text,
           'startSeconds':round(cursor,3),'durationSeconds':round(seconds,3)}
        if other:x.update(other)
        items.append(x)
        cursor=round(cursor+seconds,3)
    if module=='02':
        for seg,subtitle_text in zip(segments,strike['texts']):
            original=clean(subtitle_text)
            if not original:continue  # Omite CUT numai când tot textul este tăiat
            chunks=split_subtitles_readable(original,74)
            if not chunks or clean(' '.join(chunks))!=original:
                raise ValueError('Text de subtitrare pierdut la CUT-%02d'%seg['ordinal'])
            group='CUT-%02d'%seg['ordinal']
            style='Italic' if seg.get('reporter_italic') is True else 'Regular'
            for chunk in chunks:
                add(group,chunk,duration_for(chunk,minimum=.9,maximum=6.0,cps=13.0),
                    {'fontStyle':style,'sourceParagraph':seg['paragraph_start']})
    elif module=='03':
        markers=_editorial_markers(docx)
        recorded=set()
        for seg,name,block in _name_blocks(segments,markers):
            if name is None:
                kind,reason=_skip_identity_reason(seg,markers)
                skipped.append({'paragraph':seg['paragraph_start'],
                                'cutId':'CUT-%02d'%seg['ordinal'],
                                'sourceCode':seg.get('source_code'),
                                'kind':kind,'reason':reason})
                continue
            if block in recorded:continue
            recorded.add(block)
            person,title,anonymous=name
            add('BLOCK-'+str(block),person,duration_for(person+' '+title,minimum=3.3,maximum=5.5,cps=16),
                {'name':person,'title':title,'anonymous':bool(anonymous),'sourceParagraph':seg['paragraph_start']})
    else:
        if not captions:raise ValueError('Nu există cartoane editoriale în DOCX')
        for j,block in enumerate(captions):
            original=clean(block['text'])
            pages=build_pages(original)
            if not pages or clean(' '.join(page['text'] for page in pages))!=original:
                raise ValueError('Textul cartonului a fost modificat: bloc '+str(j+1))
            for k,page in enumerate(pages):
                raw=page['text']; slots=page['physicalLines']; bold=page['boldWords']
                add('VO-'+str(j+1),raw,duration_for(raw,minimum=3,maximum=9,cps=13.5),
                    {'physicalLines':slots,'textLine0':slots[0],'textLines123':'\r'.join(slots[1:]),
                     'boldWords':bold,'emphasisPhrases':page['emphasisPhrases'],
                     'colorScheme':'Orange','sourceParagraph':block['paragraph'],
                     'page':k+1,'pageCount':len(pages)})
    if not items:raise ValueError('Nu există elemente valide pentru categoria '+module)
    # Unele sincronuri nu au IDENTIFICARE editorială: nu inventa nume și
    # nu opri burtierele valide. Raportează exact fiecare CUT nerecunoscut,
    # iar CEP cere acceptare explicită ÎNAINTE de orice import pe V4.
    if module=='03':
        for entry in skipped:
            if not isinstance(entry.get('cutId'),str) or not re.fullmatch(r'CUT-\d+',entry['cutId']):
                raise ValueError('Identificare invalidă în lista REVIEW.')
            if not isinstance(entry.get('paragraph'),int) or entry['paragraph']<0:
                raise ValueError('Paragraf invalid în lista REVIEW.')
    review_required=any(entry.get('kind')=='REVIEW_MISSING_IDENTITY' for entry in skipped)
    return {'schema':SCHEMA,'module':module,'mode':'independent',
            'timing':'SCRIPT_PREVIEW' if module=='02' else 'EDITORIAL_PREVIEW',
            'docx_sha256':hashlib.sha256(docx.read_bytes()).hexdigest(),
            'sourceDocx':docx.name,'separatorSeconds':separator,'items':items,
            'totalItems':len(items),'groups':len({a['groupId'] for a in items}),
            'skipped':skipped,'reviewRequired':review_required,
            'strikePolicy':strike['policy'] if strike else 'NOT_APPLICABLE',
            'omittedStrikeCuts':strike['cutIdsOmitted'] if strike else [],
            'removedStrikeCharacters':strike['removedCharacters'] if strike else 0,
            'strikeCutsChanged':strike['cutsChanged'] if strike else 0,
            'strikeAudit':strike_audit,
            'strikeAnalysisReady':True if module=='02' else False,
            'reviewPolicy':('EXPLICIT_SKIP_APPROVAL' if review_required else
                            'AUTO_AMBIENCE_SKIP' if skipped else 'NONE'),
            'containsAudioSync':False,'generatorVersion':'R2.1.1-auto-preflight-readable'}

def main():
    p=argparse.ArgumentParser()
    p.add_argument('--docx',required=True,type=Path);p.add_argument('--module',required=True,choices=('02','03','04'))
    p.add_argument('--output',required=True,type=Path);p.add_argument('--separator',default=.75,type=float)
    p.add_argument('--omit-struck',action='store_true',help='Omit only Word strikethrough text from module 02')
    x=p.parse_args()
    plan=build(x.docx,x.module,x.separator,omit_struck=x.omit_struck)
    x.output.parent.mkdir(parents=True,exist_ok=True)
    tmp=x.output.with_suffix('.json.tmp')
    tmp.write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')
    tmp.replace(x.output)
    print(json.dumps({'ok':True,'plan':str(x.output),'count':plan['totalItems'],'groups':plan['groups'],
                      'docx_sha256':plan['docx_sha256'],'timing':plan['timing'],
                      'skipped':plan['skipped'],'reviewRequired':plan['reviewRequired'],
                      'strikePolicy':plan['strikePolicy'],
                      'omittedStrikeCuts':plan['omittedStrikeCuts'],
                      'removedStrikeCharacters':plan['removedStrikeCharacters'],
                      'strikeCutsChanged':plan['strikeCutsChanged'],
                      'strikeAuditCount':len(plan['strikeAudit']),
                      'strikeAnalysisReady':plan['strikeAnalysisReady'],
                      'generatorVersion':plan['generatorVersion']},ensure_ascii=False))
if __name__=='__main__':main()
