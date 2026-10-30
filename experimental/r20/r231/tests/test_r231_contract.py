"""R231 acceptance: a new DOCX can independently create 02, 03, 04 plans
without 01 Montage and without ENGINE_VERSION_MISMATCH.

Built with the *published R2.1.2* workflow + *R2.3.1* SHA-locked parser.
All user-source formats are represented by synthetic data; private docs
are tested locally, never uploaded to GitHub.
"""
from __future__ import annotations
import contextlib,hashlib,io,json,runpy,sys,tempfile
from pathlib import Path
from docx import Document
# Import R230 integration synthetic dependency adapters, executing its existing
# eight tests as well. This is the same original R2.1.2 wrapper.
runpy.run_path(str(Path(__file__).resolve().parents[2]/"r230/tests/test_r230_integration.py"))
from r20_graphics_plans import build,main
SHA="a939e588fa4c458c05ceaae2044539564a21ba4c0975c9294089b518ed84d770"
engine=Path('/tmp/aive_runtime/tools/r230_universal.py')
assert hashlib.sha256(engine.read_bytes()).hexdigest()==SHA
allowed={'R2.1.1-auto-preflight-readable','R2.3-UNIVERSAL-DOCX-9-FORMATS'}
scenarios=[
 ("SINCRON+COD+UPPERCASE",["SINCRON: Ana Exemplu, pensionară","C9625","00:10–00:20",
   "Astăzi venim la primărie pentru ajutor.","UN PROIECT NOU PENTRU LOCUITORII SATULUI"]),
 ("VOX",["VOX: Zinaida Exemplu, pensionară","0656","00:12–00:23",
   "Sunt uimită de această situație, dar sper că vom găsi o soluție.",
   "OAMENII DIN SAT AȘTEAPTĂ O DECIZIE"]),
 ("SOT",["PKG","SOT [07:08] [07:18]","Vera Exemplu","Fiica Annei",
   "Am aflat despre situație din presă.","UN NOU MESAJ PENTRU CETĂȚENI"]),
 ("DOCUMENTARY",["CARTON:","O persoană locuiește în această comună.",
   "SINCRON: Radu Exemplu, profesor","C9623","05:09-05:15",
   "Acest proiect poate ajuta persoanele în vârstă."]),
 ("LOCATION",["Script:","Geo: Herson","Ana Exemplu, locuitoare din Herson",
   "Sperăm să fie pace în oraș.","OAMENII REVIN LA VIAȚA LOR OBIȘNUITĂ"])
]
def check(plan,report,module,doc,omit):
    assert report['ok'] and plan['mode']=='independent'
    assert plan['docx_sha256']==report['docx_sha256']==hashlib.sha256(doc.read_bytes()).hexdigest()
    assert plan['module']==module and plan['totalItems']==report['count']
    assert plan['generatorVersion']==report['generatorVersion']
    assert plan['generatorVersion'] in allowed
    assert report['strikeAnalysisReady']==plan['strikeAnalysisReady']==(module=='02')
    assert report['strikeAuditCount']==len(plan['strikeAudit'])
    assert plan['strikePolicy']==report['strikePolicy']==(
        ('OMIT_STRUCK' if omit else 'INCLUDE_STRUCK') if module=='02' else 'NOT_APPLICABLE')
    assert plan['strikeCutsChanged']==report['strikeCutsChanged']
    assert plan['removedStrikeCharacters']==report['removedStrikeCharacters']
    assert plan['omittedStrikeCuts']==report['omittedStrikeCuts']
    assert report['reviewRequired']==plan['reviewRequired']
    assert plan['reviewRequired']==any(x['kind']=='REVIEW_MISSING_IDENTITY' for x in plan['skipped'])
    assert plan['reviewPolicy']==(
        'EXPLICIT_SKIP_APPROVAL' if plan['reviewRequired'] else
        'AUTO_AMBIENCE_SKIP' if plan['skipped'] else 'NONE')
    if module=='02':
        assert all(len(x['text'])<=74 for x in plan['items'])
        assert all(x['fontStyle'] in ('Regular','Italic') for x in plan['items'])
    if module=='03':
        assert all(x.get('name') and x.get('title') for x in plan['items'])
    if module=='04':
        assert all(len(l)<=33 for x in plan['items'] for l in x['physicalLines'])
        assert all(x['colorScheme']=='Orange' for x in plan['items'])
with tempfile.TemporaryDirectory() as t:
    for title,rows in scenarios:
        doc=Path(t)/(title+'.docx')
        d=Document()
        for row in rows:d.add_paragraph(row)
        d.save(doc)
        for mod in ('02','03','04'):
            # exact module entrypoint accepts command-line params and writes JSON
            output=Path(t)/(title+'-'+mod+'.json')
            args=['r20_graphics_plans.py','--docx',str(doc),'--module',mod,'--output',str(output)]
            saved=sys.argv[:]
            try:
                sys.argv=args
                out=io.StringIO()
                with contextlib.redirect_stdout(out):main()
            finally:sys.argv=saved
            report=json.loads(out.getvalue())
            plan=json.loads(output.read_text(encoding='utf8'))
            check(plan,report,mod,doc,False)
            print('PASS DOCX selected in 00 ->',title,mod,'independent',
                'items',len(plan['items']),'ENGINE_VERSION_MISMATCH impossible')
    # A genuine Word strikethrough mark, not a forged Python source field.
    doc=Path(t)/'strike.docx'
    d=Document()
    d.add_paragraph('SINCRON: Ana Exemplu, pensionară')
    p=d.add_paragraph();p.add_run('Text păstrat. ');s=p.add_run('Text eliminat.');s.font.strike=True
    d.add_paragraph('UN CARTON EDITORIAL ÎN FORMAT VERTICAL')
    d.save(doc)
    plan=build(doc,'02',omit_struck=True)
    assert plan['strikePolicy']=='OMIT_STRUCK'
    assert plan['strikeAnalysisReady'] is True
    assert plan['strikeCutsChanged']>=1 and plan['removedStrikeCharacters']>0
    assert all('eliminat' not in x['text'] for x in plan['items'])
    print('PASS DOCX genuine Strikethrough: omitted only crossed-out words')
print('R231 VERIFIED: original R2.1.2 main+new universal -> panel schema for 02/03/04')
