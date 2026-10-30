"""Synthetic acceptance for R2.3.2: Word-run speaker style, CN table,
no fabricated NameTags and no dropped table continuation.
The examples are synthetic, not newsroom files.
"""
import hashlib,sys,tempfile,types
from pathlib import Path
from docx import Document
TOOLS=Path(__file__).resolve().parents[1]/"candidate"/"tools"
sys.path.insert(0,str(TOOLS))
# Production runner from v2.3.0 is later loaded by CI under /tmp/aive_runtime.
# Use the exact approved helpers for splitting/layout when present.
RUNTIME=Path('/tmp/aive_runtime/tools')
sys.path.insert(1,str(RUNTIME))
if not (RUNTIME/'r20_graphics_plans.py').exists():
    def split(t,limit=74):
        words=t.split();current='';out=[]
        for w in words:
            if current and len(current)+1+len(w)>limit:
                out.append(current);current=w
            else:current=(current+' '+w).strip()
        if current:out.append(current)
        return out
    sys.modules['r20_graphics_plans']=types.SimpleNamespace(split_subtitles_readable=split)
    sys.modules['r204_carton_layout']=types.SimpleNamespace(build_pages=lambda t:[])
from r230_universal import parse,build,clean

def save(doc,folder,name):
    out=Path(folder)/(name+'.docx')
    doc.save(out)
    return out
with tempfile.TemporaryDirectory() as temp:
    # The exact mixed style from the screenshot: one italic "Da."
    doc=Document()
    doc.add_paragraph('VOX')
    doc.add_paragraph('0651')
    doc.add_paragraph('00.10–00.14')
    p=doc.add_paragraph()
    p.add_run('Datele a peste 700 de mii de cetățeni au fost accesate de hackeri. ')
    p.add_run('Din Moldova aveți în vedere? ')
    em=p.add_run('Da. ');em.italic=True
    p.add_run('Nu am auzit.')
    source=save(doc,temp,'mixed_italic')
    parsed=parse(source)
    seg=parsed['segments']
    assert len(seg)==3,[(x['text'],x['reporter_italic']) for x in seg]
    assert [x['reporter_italic'] for x in seg]==[False,True,False]
    assert [x['text'] for x in seg]==[
      'Datele a peste 700 de mii de cetățeni au fost accesate de hackeri. Din Moldova aveți în vedere?',
      'Da.','Nu am auzit.']
    assert seg[1]['speaker'] is None and seg[1]['role'] is None
    subtitles=build(source,'02')
    captions=subtitles['items']
    assert len(captions)>=3
    assert any(x['text']=='Da.' and x['fontStyle']=='Italic' for x in captions)
    assert not any(x['fontStyle']=='Regular' and 'Da.' in x['text'] for x in captions)
    assert all(len(x['text'])<=74 for x in captions)
    assert clean(' '.join(x['text'] for x in captions))==clean(' '.join(x['text'] for x in seg))
    tags=None
    try:tags=build(source,'03')
    except ValueError as e:assert 'Nicio burtieră validă' in str(e)
    assert tags is None
    print('PASS mixed Word runs: Da. is a separate Italic reporter cue; Regular interview unchanged')

    # Exactly the sort of two-column narration in screenshot. The table
    # contains a row where only the right cell is populated. This used to
    # be dropped entirely by the R230 table iterator.
    doc=Document()
    t=doc.add_table(rows=4,cols=2)
    t.cell(0,0).text='Facebook - Arzu Abdullah Gul Zaman\nGEO - Jalilabad, 7 octombrie'
    t.cell(0,1).text='Fetița din imagine nu este admisă la școală pentru că poartă hijab.'
    t.cell(1,0).text=''
    t.cell(1,1).text='Este o consecință a noului cod vestimentar în școli.'
    t.cell(2,0).text='GEO - Nardaran, 12 octombrie\nFacebook - Elcin Alioglu'
    t.cell(2,1).text='Noile reguli au reaprins dezbateri despre libertatea religioasă.'
    t.cell(3,0).text='Gulmira Aslanova\nmamă'
    paragraph=t.cell(3,1).paragraphs[0]
    rr=paragraph.add_run('Articolul 48 din Constituție ne garantează libertatea religioasă.')
    rr.bold=True
    source=save(doc,temp,'two_col_CN')
    parsed=parse(source)
    segments=parsed['segments']
    assert len(segments)==4,[(x['text'],x['kind']) for x in segments]
    assert [x['reporter_italic'] for x in segments]==[True,True,True,False]
    assert [x['kind'] for x in segments]==['narration']*3+['table_sot']
    assert segments[-1]['speaker']=='Gulmira Aslanova'
    assert segments[-1]['role']=='mamă'
    captions=build(source,'02')['items']
    assert any('Fetița din imagine' in x['text'] and x['fontStyle']=='Italic' for x in captions)
    assert any('consecință a noului cod vestimentar' in x['text'] and x['fontStyle']=='Italic' for x in captions)
    assert any('Articolul 48' in x['text'] and x['fontStyle']=='Regular' for x in captions)
    nametags=build(source,'03')
    assert len(nametags['items'])==1
    assert nametags['items'][0]['name']=='Gulmira Aslanova'
    assert nametags['items'][0]['title']=='mamă'
    assert all(x['kind']=='SKIP_NARRATION' for x in nametags['skipped'])
    assert nametags['reviewRequired'] is False
    assert nametags['reviewPolicy']=='AUTO_AMBIENCE_SKIP'
    print('PASS CN in Word two-column table, including right-only continuation; Italic CN, Regular SOT and one NameTag')

    # Speaker with an italic reporter interruption in the same paragraph:
    doc=Document()
    p=doc.add_paragraph('SINCRON: Ana Exemplu, pensionară')
    p=doc.add_paragraph()
    p.add_run('O declarație. ')
    em=p.add_run('Dar de ce? ');em.italic=True
    p.add_run('Așa cred eu.')
    source=save(doc,temp,'intra_speaker')
    seg=parse(source)['segments']
    assert [x['speaker'] for x in seg]==['Ana Exemplu',None,'Ana Exemplu']
    assert [x['reporter_italic'] for x in seg]==[False,True,False]
    tag=build(source,'03')
    assert all(item['name']=='Ana Exemplu' for item in tag['items'])
    assert not any(i.get('kind')=='REVIEW_MISSING_IDENTITY' for i in tag['skipped'])
    print('PASS reporter turn never inherits previous speaker identity and cannot create a fake NameTag')
print('R232 parser acceptance complete')
