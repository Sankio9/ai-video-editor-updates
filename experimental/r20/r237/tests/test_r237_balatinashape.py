"""R2.3.7 Balatina *structural* fixture (synthetic text, not user's words).
Ensures 8 timed editorial paragraphs incl 0012 ambient, 8 uppercase
paragraphs in 5 VO groups and six reintroduced identity blocks.
"""
import sys,tempfile,types
from pathlib import Path
from docx import Document
ROOT=Path('/tmp/r236_runtime/tools');sys.path.insert(0,str(ROOT))
def layout(text):
    lines=[];cur=''
    for word in text.split():
        if cur and len(cur)+len(word)+1>28:
            lines.append(cur);cur=word
        else:cur=(cur+' '+word).strip()
    if cur:lines.append(cur)
    result=[]
    for i in range(0,len(lines),4):
        slots=(lines[i:i+4]+['','','',''])[:4]
        result.append({'text':' '.join(lines[i:i+4]),'physicalLines':slots,
                       'boldWords':{},'emphasisPhrases':[]})
    return result
def split(txt,limit=74):
    cues=[];cur=''
    for w in txt.split():
        if cur and len(cur)+len(w)+1>limit:cues.append(cur);cur=w
        else:cur=(cur+' '+w).strip()
    if cur:cues.append(cur)
    return cues
def clean(s):return ' '.join(s.split())
sys.modules["reportaj_r1_parse_docx"]=types.SimpleNamespace(extract=lambda x:{"segments":[],"graphics":[]})
sys.modules["r193_docx_reader"]=types.SimpleNamespace(read=lambda x:{"segments":[],"captions":[]})
sys.modules["r161_prepare_nametags"]=types.SimpleNamespace(card_for=lambda x:None)
sys.modules["r171_layout"]=types.SimpleNamespace(clean=clean)
sys.modules["r204_carton_layout"]=types.SimpleNamespace(build_pages=layout)
sys.modules["r207_docx_strike"]=types.SimpleNamespace(extract_strike_policy=lambda *args:None)
sys.modules["r153_prepare"]=types.SimpleNamespace(split_balanced=lambda x,limit=74:split(x,limit))
from r230_universal import parse,build as universal_build
from r237_coverage import eligible,verify
from r20_graphics_plans import build as canonical_build

with tempfile.TemporaryDirectory() as folder:
    p=Path(folder)/'realistic editorial format.docx'
    d=Document()
    d.add_paragraph('Titlu document jurnalistic despre profesie')
    d.add_paragraph('Introducerea documentului nu devine subtitrare.')
    d.add_paragraph('Autor One, Autor Two')
    d.add_paragraph('Ambianță')
    d.add_paragraph('0012')
    d.add_paragraph('00.09 O intervenție despre situația în comunitate este redată aici. 00.20')
    doc_vo_groups=[
       ['PERSOANA ESTE PREZENTĂ ZILNIC ÎN COMUNITATEA DIN SATUL EXEMPLU',
        'A VENIT SĂ MUNCEASCĂ DUPĂ STUDIILE DE SPECIALITATE'],
       ['TOȚI PROFESIONIȘTII CONTINUĂ SĂ AJUTE OAMENII DIN ACEST RAION'],
       ['MEDICII DIN SAT LUCREAZĂ PENTRU PACIENȚII LOR ÎN FIECARE ZI'],
       ['CENTRUL MEDICAL OFERĂ SERVICII PENTRU SATUL EXEMPLU',
        'PACIENȚII AU ACCES LA INVESTIGAȚII ȘI CONSULTAȚII'],
       ['ÎN COMUNITATE SUNT MAI MULȚI PROFESIONIȘTI CU EXPERIENȚĂ',
        'SISTEMUL ARE NEVOIE DE PERSONAL SUPLIMENTAR ÎN SATE']
    ]
    blocks=[('0007','Nina Exemplu','Șefa Centrului de Sănătate'),
            ('0007','Nina Exemplu','Șefa Centrului de Sănătate'),
            ('0034','Eugenia Exemplu','Medic de familie'),
            ('0034','Eugenia Exemplu','Medic de familie'),
            ('0007','Nina Exemplu','Șefa Centrului de Sănătate'),
            ('0007','Nina Exemplu','Șefa Centrului de Sănătate')]
    def marker(code,name,role,content,allstruck=False):
        d.add_paragraph(code)
        d.add_paragraph(name)
        d.add_paragraph(role)
        q=d.add_paragraph()
        if allstruck:
            r=q.add_run(content)
            r.font.strike=True
        else:
            q.add_run(content)
    for part in doc_vo_groups[0]:d.add_paragraph(part)
    marker(*blocks[0],'1.32 O declarație despre munca în comunitate exprimă opinia ei. 1.51')
    for part in doc_vo_groups[1]:d.add_paragraph(part)
    marker(*blocks[1],'9.00 Un răspuns despre infrastructura rurală explică greutățile zilnice. 9.38')
    marker(*blocks[2],'4.01 Persoana descrie colectivul și relațiile de la serviciu. 4.06')
    for part in doc_vo_groups[2]:d.add_paragraph(part)
    marker(*blocks[3],'4.21 O altă declarație explică experiența de lucru de-a lungul anilor. 4.48')
    paragraph=d.add_paragraph()
    r=paragraph.add_run('5.10 Această frază este integral exclusă prin Strikethrough și nu se afișează. 5.26')
    r.font.strike=True
    for part in doc_vo_groups[3]:d.add_paragraph(part)
    marker(*blocks[4],'10.41 Persoana descrie schimbările observate în comunitate. 11.12')
    for part in doc_vo_groups[4]:d.add_paragraph(part)
    marker(*blocks[5],'4.28 Ultima replică explică experiența de viață și munca în sat. 5.00')
    d.save(p)
    audio,vo=eligible(p)
    assert len(audio)==8,(len(audio),[x['idx'] for x in audio])
    assert len(vo)==8,(len(vo),[x['idx'] for x in vo])
    parsed=parse(p)
    audit=verify(p,parsed,omit_struck=True,module='02')
    assert audit['timedParagraphs']==8 and audit['voParagraphs']==8,audit
    assert len(parsed['segments'])==8,len(parsed['segments'])
    assert len(parsed['captions'])>=5,len(parsed['captions'])
    v3=canonical_build(p,'02',omit_struck=True)
    assert v3['editorialCoverage']['timedParagraphs']==8
    assert 'intervenție despre situația' in v3['items'][0]['text']
    assert not any('Această frază este integral exclusă' in x['text'] for x in v3['items'])
    assert v3['omittedStrikeCuts'],v3['strikeAudit']
    v4=canonical_build(p,'03')
    assert v4['totalItems']==6,v4['items']
    v5=canonical_build(p,'04')
    assert v5['editorialCoverage']['verifiedVoParagraphs']==8
    all_text=' '.join(x['text'] for x in v5['items'])
    for block in doc_vo_groups:
        for text in block:
            assert text in all_text,text
    assert all(len(line)<=28 for x in v5['items'] for line in x['physicalLines'])
    print('PASS private-document structure: 8 timed, 8 uppercase in 5 groups, 6 NameTags, first Ambianță present, full-strike excluded')
    print('PASS all Word uppercase VO paragraphs retained, max 28 chars per physical line in fixture')
