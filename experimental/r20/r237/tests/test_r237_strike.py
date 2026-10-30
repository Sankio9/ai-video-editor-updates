"""R237 optional montage Strikethrough proof: temporary Word only, immutable
original SHA and per-paragraph audit. No source asset modifications.
"""
from __future__ import annotations
import sys,tempfile,unittest
from pathlib import Path
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'tools'))
from r237_strikethrough_docx import filter_docx,verify
class Tests(unittest.TestCase):
 def test_mixed_regular_and_struck_docx_is_filtered_privately(self):
  with tempfile.TemporaryDirectory() as tmp:
   folder=Path(tmp)
   path=folder/'editorial.docx'
   d=Document();d.add_paragraph('SINCRON');d.add_paragraph('0034')
   p=d.add_paragraph()
   p.add_run('Păstrăm aceste cuvinte, ')
   r=p.add_run('eliminăm formularea greșită, ');r.font.strike=True
   p.add_run('iar finalul rămâne.')
   d.save(path)
   orig=path.read_bytes()
   r=filter_docx(path,folder/'audit')
   self.assertGreater(r['removedCharacters'],0)
   self.assertEqual(path.read_bytes(),orig)
   copy=Document(r['analysisDOCX'])
   self.assertEqual(copy.paragraphs[1].text,'0034')
   self.assertEqual(copy.paragraphs[2].text,'Păstrăm aceste cuvinte, iar finalul rămâne.')
   self.assertEqual(r['changedParagraphs'],1)
   manifest=folder/'audit'/('r237_strike_'+r['originalSHA256'][:24]+'.json')
   assert verify(path,Path(r['analysisDOCX']),manifest)['filteredSHA256']==r['filteredSHA256']
   copy_par=copy.paragraphs[2].text
   self.assertNotIn('greșită',copy_par)
   print('PASS Word mixed runs, original untouched, private filtered DOCX + SHA audit')
 def test_complete_struck_paragraph_must_stop(self):
  with tempfile.TemporaryDirectory() as tmp:
   path=Path(tmp)/'word.docx';d=Document()
   p=d.add_paragraph();r=p.add_run('Tot acest sincron este tăiat.');r.font.strike=True
   d.save(path)
   with self.assertRaisesRegex(ValueError,'STRUCK_ENTIRE_PARAGRAPH'):
    filter_docx(path,Path(tmp)/'temp')
   self.assertEqual(Document(path).paragraphs[0].text,'Tot acest sincron este tăiat.')
   print('PASS all-struck paragraph cannot silently change CUT count')
 def test_struck_source_code_stops(self):
  with tempfile.TemporaryDirectory() as tmp:
   path=Path(tmp)/'word.docx';d=Document()
   p=d.add_paragraph();r=p.add_run('0034');r.font.strike=True
   p.add_run(' Cineva')
   d.save(path)
   with self.assertRaisesRegex(ValueError,'STRUCK_SOURCE_CODE'):
    filter_docx(path,Path(tmp)/'temp')
   print('PASS source suffix cannot silently be deleted')
 def test_provenance_detects_source_docx_edit(self):
  with tempfile.TemporaryDirectory() as tmp:
   path=Path(tmp)/'word.docx';d=Document()
   p=d.add_paragraph();p.add_run('Bună ');r=p.add_run('inutil ');r.font.strike=True;p.add_run('ziua.')
   d.save(path)
   report=filter_docx(path,Path(tmp)/'temp')
   manifest=Path(tmp)/'temp'/('r237_strike_'+report['originalSHA256'][:24]+'.json')
   new=Document(path);new.add_paragraph('Text introdus după analiză.');new.save(path)
   with self.assertRaisesRegex(ValueError,'STRIKE_PROVENANCE_CHANGED'):
    verify(path,report['analysisDOCX'],manifest)
   print('PASS changed original Word fails immutable provenance check')
if __name__=='__main__':unittest.main(verbosity=2)
