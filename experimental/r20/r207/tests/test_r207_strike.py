"""R2.0.7: synthetic DOCX formatting regressions, no real script/public data."""
from __future__ import annotations
import tempfile,sys,unittest
from pathlib import Path
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/"tools"))
from r207_docx_strike import extract_strike_policy

class TestStrikePolicy(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.path=Path(self.tmp.name)/"synthetic.docx"
        self.doc=Document()
    def tearDown(self):self.tmp.cleanup()
    def put(self,parts,kind="normal"):
        p=self.doc.add_paragraph()
        for t,mode in parts:
            r=p.add_run(t)
            if mode=="strike":r.font.strike=True
            if mode=="double":r.font.double_strike=True
            if mode=="italic":r.italic=True
            if mode=="bold":r.bold=True
        return p
    def build(self,segments):
        self.doc.save(self.path)
        return extract_strike_policy(self.path,segments,True)
    @staticmethod
    def cut(i,para,text):
        return {"ordinal":i,"paragraph_start":para,"text":text}

    def test_exact_strike_excluded_italic_preserved(self):
        p=self.put([
            ("Vreau să mă dezvolt în profesia dată. ","strike"),
            ("Unde ai vrea să muncești? ","italic"),
            ("Să fiu bucătar în Anglia.","normal")])
        s=p.text
        a=self.build([self.cut(2,0,s)])
        self.assertEqual(a["texts"],["Unde ai vrea să muncești? Să fiu bucătar în Anglia."])
        self.assertEqual(a["removedCharacters"],len("Vreausămădezvoltînprofesiadată."))
        self.assertEqual(a["cutsChanged"],1)
        self.assertEqual(a["cutIdsOmitted"],[])
        self.assertEqual(a["policy"],"OMIT_STRUCK")
    def test_default_include_preserves_original(self):
        self.put([("Omite acest text","strike")])
        self.doc.save(self.path)
        x=extract_strike_policy(self.path,[self.cut(1,0,"Omite acest text")],False)
        self.assertEqual(x["texts"],["Omite acest text"])
        self.assertEqual(x["removedCharacters"],0)
        self.assertEqual(x["policy"],"INCLUDE_STRUCK")
    def test_italic_not_removed(self):
        p=self.put([("Reporter: ","italic"),("Ce mai faci?","normal")])
        self.assertEqual(self.build([self.cut(1,0,p.text)])["texts"],[p.text])
    def test_double_strike(self):
        p=self.put([("NU ","double"),("DA","bold")])
        self.assertEqual(self.build([self.cut(1,0,p.text)])["texts"],["DA"])
    def test_partial_word_and_punctuation(self):
        p=self.put([("bu","strike"),("cătar!", "normal")])
        self.assertEqual(self.build([self.cut(1,0,p.text)])["texts"],["cătar!"])
    def test_entire_cut_omitted(self):
        p=self.put([("Scos complet", "strike")])
        x=self.build([self.cut(3,0,p.text)])
        self.assertEqual(x["texts"],[""])
        self.assertEqual(x["cutIdsOmitted"],["CUT-03"])
    def test_header_never_omitted_by_context(self):
        self.put([("SINCRON – 9981", "strike")])
        p=self.put([("Textul valid","normal")])
        self.assertEqual(self.build([self.cut(1,0,p.text)])["texts"],["Textul valid"])
    def test_two_cuts_separate(self):
        a=self.put([("Partea tăiată ","strike"),("Partea bună","normal")])
        b=self.put([("Text simplu","normal")])
        x=self.build([self.cut(1,0,a.text),self.cut(2,1,b.text)])
        self.assertEqual(x["texts"],["Partea bună","Text simplu"])
        self.assertEqual(x["cutsChanged"],1)
    def test_spanning_multiple_paragraphs(self):
        self.put([("Nu se include","strike")])
        self.put([("Dar păstrăm acest text","normal")])
        x=self.build([self.cut(1,0,"Nu se include Dar păstrăm acest text")])
        self.assertEqual(x["texts"],["Dar păstrăm acest text"])
    def test_no_guess_on_mismatch(self):
        self.put([("Altă formulare","normal")])
        with self.assertRaisesRegex(ValueError,"STRIKE_FORMAT_AMBIGUOUS"):
            self.build([self.cut(1,0,"Acesta nu există")])
    def test_identical_duplicate_in_one_window_is_blocked(self):
        self.put([("REPLICĂ REPLICĂ","normal")])
        with self.assertRaisesRegex(ValueError,"STRIKE_FORMAT_AMBIGUOUS"):
            self.build([self.cut(1,0,"REPLICĂ")])
    def test_no_cross_source_silently(self):
        self.put([("Text 1","normal")]);self.put([("Text 2","normal")])
        with self.assertRaisesRegex(ValueError,"STRIKE_FORMAT_AMBIGUOUS"):
            self.build([self.cut(1,0,"Text 2"),self.cut(2,1,"Text 2")])
    def test_invalid_ordinal_paragraph(self):
        self.put([("Text","normal")])
        with self.assertRaisesRegex(ValueError,"paragraph_start"):
            self.build([self.cut(1,-1,"Text")])
    def test_explicit_boolean_only(self):
        self.put([("Text","normal")]);self.doc.save(self.path)
        with self.assertRaisesRegex(ValueError,"bool"):
            extract_strike_policy(self.path,[self.cut(1,0,"Text")],"yes")
    def test_empty_target_does_not_mutate(self):
        self.put([("XYZ","normal")])
        x=self.build([self.cut(1,0,"")])
        self.assertEqual(x["texts"],[""])
        self.assertEqual(x["cutIdsOmitted"],[])
    def test_underlined_and_bold_text_remains(self):
        p=self.doc.add_paragraph()
        p.add_run("Text ").bold=True
        p.add_run("păstrat").underline=True
        self.assertEqual(self.build([self.cut(1,0,p.text)])["texts"],["Text păstrat"])

if __name__=="__main__":unittest.main(verbosity=2)
