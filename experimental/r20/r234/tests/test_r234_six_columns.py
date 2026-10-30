"""Regressions reproducing real 2026-10-15 newsroom Word structure:
SIX physical DOCX grid columns; merged source 0-1; merged editorial
text 2-3; empty 4-5; INTRO then PKG; right-only CN; 3 named SOT.
All sentences in this test are synthetic. Uploaded newsroom DOCX stays private.
"""
from pathlib import Path
from docx import Document
import sys,tempfile,types
root=Path(__file__).resolve().parents[1]/"candidate"/"tools"
sys.path.insert(0,str(root))
def splitter(text,limit=74):
 words=text.split();parts=[];line=''
 for w in words:
  if line and len(line)+1+len(w)>limit:
   parts.append(line);line=w
  else:line=(line+' '+w).strip()
 if line:parts.append(line)
 return parts
sys.modules["r20_graphics_plans"]=types.SimpleNamespace(split_subtitles_readable=splitter)
def pages(text):
 lines=[];line=''
 for w in text.split():
  if line and len(line)+1+len(w)>33:
   lines.append(line);line=w
  else:line=(line+' '+w).strip()
 if line:lines.append(line)
 return [{"text":" ".join(lines[i:i+4]),
          "physicalLines":(lines[i:i+4]+["","","",""])[:4],
          "boldWords":{"0":"","1":"","2":"","3":""},
          "emphasisPhrases":[]} for i in range(0,len(lines),4)]
sys.modules["r204_carton_layout"]=types.SimpleNamespace(build_pages=pages)
from r230_universal import parse,build,paragraph_rows
def merged_line(table,i,left="",right="",bold=False):
 r=table.rows[i]
 c=r.cells[0].merge(r.cells[1]);c.text=left
 v=r.cells[2].merge(r.cells[3]);v.text=right
 if bold:
  for p in v.paragraphs:
   for run in p.runs:run.bold=True
with tempfile.TemporaryDirectory() as folder:
 p=Path(folder)/"2026.10.15 structură sintetică.docx"
 doc=Document()
 doc.add_paragraph("Un titlu jurnalistic, care NU trebuie transformat în subtitrare.")
 table=doc.add_table(rows=11,cols=6)
 merged_line(table,0,"INTRO")
 merged_line(table,1,"Titlul știrii","Un sumar introductiv, care NU face parte din reportaj.")
 merged_line(table,2,"PKG")
 merged_line(table,3,"Facebook - autor\nGEO - localitate, 7 octombrie",
             "Prima propoziție a naratorului oferă context.\nA doua propoziție continuă contextul.")
 merged_line(table,4,"","Continuarea naratorului trebuie păstrată chiar dacă stânga este goală.")
 merged_line(table,5,"Gulmira Exemplu\nmamă",
             "Primul sincron declarat aparține unei persoane identificate.",True)
 merged_line(table,6,"RFE/RL",
             "După primul sincron, naratorul explică în continuare situația.")
 merged_line(table,7,"Elnara Exemplu\nDeputată, Partid",
             "O a doua persoană oferă o declarație în această secvență.",True)
 merged_line(table,8,"Sanubar Exemplu\nSociolog",
             "Al treilea sincron include un răspuns al unui expert.",True)
 merged_line(table,9,"",
             "Fraza de încheiere este narator, chiar dacă sursa lipsește.")
 merged_line(table,10,"Reuters","Reporterul explică semnificația declarațiilor.")
 doc.save(p)
 parsed=parse(p)
 seg=parsed["segments"]
 assert len(seg)==9,[(x["kind"],x["text"][:35]) for x in seg]
 assert [x["reporter_italic"] for x in seg]==[True,True,True,False,True,False,False,True,True]
 assert all("sumar introductiv" not in x["text"] for x in seg)
 assert all("Un titlu jurnalistic" not in x["text"] for x in seg)
 tags=[(x["speaker"],x["role"]) for x in seg if x["speaker"] and x["role"]]
 assert tags==[("Gulmira Exemplu","mamă"),("Elnara Exemplu","Deputată, Partid"),("Sanubar Exemplu","Sociolog")],tags
 for mod in ("02","03","04"):
  result=build(p,mod)
  assert result["generatorVersion"]=="R2.3-UNIVERSAL-DOCX-9-FORMATS"
  if mod=="02":
   assert result["totalItems"]>=9
   assert any(x["fontStyle"]=="Italic" and x["text"].startswith("Continuarea naratorului") for x in result["items"])
   assert any(x["fontStyle"]=="Regular" and "sincron" in x["text"] for x in result["items"])
   assert all(len(x["text"])<=74 for x in result["items"])
  elif mod=="03":
   assert result["totalItems"]==3, result["items"]
   assert all(x["name"] in ("Gulmira Exemplu","Elnara Exemplu","Sanubar Exemplu") for x in result["items"])
  else:
   assert result["totalItems"]==0
  print("PASS merged 6-col independent",mod,"items",result["totalItems"])
 rows=paragraph_rows(Document(p))
 assert sum(1 for r in rows if r.get("table"))==9,[(x.get("left"),x.get("text")[:20]) for x in rows if x.get("table")]
 print("PASS 6-col PKG skip INTRO, 9 editorial paragraphs and exact 3 verified speakers")
 # Two real grid columns, without PKG, must remain compatible.
 q=Path(folder)/"simple two columns.docx"
 d=Document();t=d.add_table(rows=2,cols=2)
 t.cell(0,0).text="Mara Exemplu\nmamă"
 t.cell(0,1).text="Aceasta este o declarație exprimată într-un interviu."
 t.cell(1,0).text="RFE/RL"
 t.cell(1,1).text="Naratorul descrie acum ceea ce s-a întâmplat."
 d.save(q)
 assert len(parse(q)["segments"])==2
 assert build(q,"03")["totalItems"]==1
 print("PASS conventional two-column DOCX still works")
print("R234 REALISTIC MERGED TABLE REGRESSION OK")
