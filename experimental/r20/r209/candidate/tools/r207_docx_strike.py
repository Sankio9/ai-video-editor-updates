"""R2.0.9 — robust opt-in DOCX strikethrough removal for subtitle text ONLY.

Never modifies the DOCX or media; never touches Montage 01, NameTag 03, Cartoane 04.
Uses runs (including hyperlink runs), preserving whitespace, punctuation and italic
text. If a parsed dialogue cannot be mapped *unambiguously* back to formatted
DOCX characters, it fails closed rather than guessing which words were crossed out.
"""
from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

from docx import Document

from docx.oxml.ns import qn

_TRUE = frozenset(("1","true","on","yes"))
_FALSE = frozenset(("0","false","off","no"))

def _xml_on_off(element):
    raw=element.get(qn("w:val"))
    if raw is None:return True
    value=str(raw).strip().lower()
    if value in _TRUE:return True
    if value in _FALSE:return False
    raise ValueError("STRIKE_FORMAT_UNSUPPORTED: invalid Word on/off property")

def _from_rpr(rpr):
    """A strike or double-strike TRUE wins within the same formatting layer."""
    if rpr is None:return None
    settings=[]
    for tag in ("strike","dstrike"):
        element=rpr.find(qn("w:"+tag))
        if element is not None:settings.append(_xml_on_off(element))
    return any(settings) if settings else None

def _style_chain(style):
    seen=set()
    for _ in range(32):
        if style is None:return None
        name=getattr(style,"style_id",None)
        if name in seen:raise ValueError("STRIKE_FORMAT_UNSUPPORTED: style inheritance cycle")
        seen.add(name)
        value=_from_rpr(style.element.find(qn("w:rPr")))
        if value is not None:return value
        style=getattr(style,"base_style",None)
    raise ValueError("STRIKE_FORMAT_UNSUPPORTED: style inheritance too deep")

def _formatting_strike(run_xml, paragraph):
    direct=run_xml.find(qn("w:rPr"))
    decision=_from_rpr(direct)
    if decision is not None:return decision

    style_ref=direct.find(qn("w:rStyle")) if direct is not None else None
    if style_ref is not None:
        style_id=style_ref.get(qn("w:val"))
        if not style_id:
            raise ValueError("STRIKE_FORMAT_UNSUPPORTED: empty run style")
        try:style=paragraph.part.styles[style_id]
        except KeyError as exc:raise ValueError("STRIKE_FORMAT_UNSUPPORTED: unknown run style") from exc
        decision=_style_chain(style)
        if decision is not None:return decision

    decision=_style_chain(paragraph.style)
    if decision is not None:return decision

    # The Word document's default run style applies even if neither paragraph
    # nor character style declares the property.
    styles=paragraph.part.styles.element
    defaults=styles.find(qn("w:docDefaults"))
    if defaults is not None:
        run_default=defaults.find(qn("w:rPrDefault"))
        if run_default is not None:
            decision=_from_rpr(run_default.find(qn("w:rPr")))
            if decision is not None:return decision
    return False

def _ancestor_deleted(run_xml, paragraph):
    node=run_xml.getparent()
    while node is not None and node is not paragraph._p:
        if node.tag in (qn("w:del"),qn("w:moveFrom")):
            return True
        node=node.getparent()
    return False

def _run_pairs(paragraph):
    """Word XML run mapping: direct, inherited strike + tracked deletions.

    lxml traverses ordinary and hyperlink runs. Deleted w:delText exists in
    Word's visual tracked-changes markup but python-docx p.text normally skips
    it; retaining provenance allows exact matching when the editorial parser
    includes the crossed-out text.
    """
    pairs=[]
    for xml_run in paragraph._p.iter(qn("w:r")):
        crossed=_ancestor_deleted(xml_run,paragraph)
        struck=crossed or _formatting_strike(xml_run,paragraph)
        for child in xml_run:
            if child.tag in (qn("w:t"),qn("w:delText")):
                text=child.text or ""
                pairs.extend((char,struck or child.tag==qn("w:delText")) for char in text)
            elif child.tag == qn("w:tab"):
                pairs.append(("\t",struck))
            elif child.tag in (qn("w:br"),qn("w:cr")):
                pairs.append(("\n",struck))
    all_text="".join(char for char,_ in pairs)
    # Direct Word p.text may omit tracked deletions. Verify the difference is
    # explained by revision markup; never assume arbitrary missing runs are safe.
    if all_text!=paragraph.text:
        visible=[]
        for xml_run in paragraph._p.iter(qn("w:r")):
            if _ancestor_deleted(xml_run,paragraph):continue
            for child in xml_run:
                if child.tag==qn("w:t"):visible.append(child.text or "")
                elif child.tag==qn("w:tab"):visible.append("\t")
                elif child.tag in (qn("w:br"),qn("w:cr")):visible.append("\n")
        if "".join(visible)!=paragraph.text:
            raise ValueError("STRIKE_FORMAT_UNSUPPORTED: Word text cannot be mapped to actual runs")
    return pairs


def _squash(pairs):
    """Whitespace-normalized text with per-character original strike provenance."""
    out=[];flags=[]
    for char,struck in pairs:
        if char.isspace():
            if out and out[-1]!=" ":
                out.append(" ");flags.append(False)
        else:
            out.append(char);flags.append(struck)
    if out and out[-1]==" ":
        out.pop();flags.pop()
    return "".join(out),flags


def _window(doc, first, last):
    if not isinstance(first,int) or not isinstance(last,int) or first<0 or last<=first or last>len(doc.paragraphs):
        raise ValueError("Invalid DOCX paragraph window")
    pairs=[]
    for para in doc.paragraphs[first:last]:
        if pairs:
            pairs.append((" ",False))
        pairs.extend(_run_pairs(para))
    return _squash(pairs)


def extract_strike_policy(docx, segments, omit_struck=False):
    """Return a list parallel to parser segments and transparent exclusion counts.

    segment 'text' is the authoritative editorial text. The original text is
    unchanged when option is off. When on, locate exact normalized text in its
    paragraph window (to the next segment header), and suppress only chars
    explicitly formatted as struck. If all dialogue was struck, omit the cue
    and record the CUT. No whole-word guessing.
    """
    if type(omit_struck) is not bool:
        raise ValueError("omit_struck must be explicit bool")
    if not isinstance(segments,list):
        raise ValueError("Invalid DOCX segments")
    if not omit_struck:
        return {"texts":[str(s["text"]) for s in segments],
                "cutIdsOmitted":[],"removedCharacters":0,"cutsChanged":0,
                "policy":"INCLUDE_STRUCK"}
    document=Document(str(docx))
    texts=[];all_omitted=[];removed=0;changed=0
    for i,seg in enumerate(segments):
        start=seg.get("paragraph_start")
        if not isinstance(start,int) or start<0 or start>=len(document.paragraphs):
            raise ValueError("CUT-%02d has invalid paragraph_start"%(i+1))
        next_start=(segments[i+1].get("paragraph_start")
                    if i+1<len(segments) else len(document.paragraphs))
        # A single paragraph may contain consecutive editorial headers/voices.
        if not isinstance(next_start,int) or next_start<start:
            raise ValueError("CUT-%02d has invalid paragraph ordering"%(i+1))
        end=max(start+1,next_start)
        if end>len(document.paragraphs):raise ValueError("Out-of-range DOCX paragraph")
        source,flags=_window(document,start,end)
        target=" ".join(str(seg["text"]).split())
        if not target:
            texts.append("");continue
        # If editorial text is repeated in a segment's source range with
        # different formatting, do not pick one at random.
        positions=[m.start() for m in re.finditer(re.escape(target),source)]
        if len(positions)!=1:
            raise ValueError("CUT-%02d: STRIKE_FORMAT_AMBIGUOUS (%d exact matches). "
                             "Cannot omit text safely."%(int(seg.get("ordinal",i+1)),len(positions)))
        offset=positions[0];region=source[offset:offset+len(target)]
        strike=flags[offset:offset+len(target)]
        remaining="".join(ch for ch,st in zip(region,strike) if not st)
        clean=" ".join(remaining.split())
        if region != clean:
            removed_chars=sum(1 for ch,st in zip(region,strike) if st and not ch.isspace())
            removed+=removed_chars
            if removed_chars:changed+=1
        if not clean:
            all_omitted.append("CUT-%02d"%int(seg.get("ordinal",i+1)))
        texts.append(clean)
    if removed == 0:
        # Never report READBACK PASS with 0 removals when the user explicitly
        # asked for struck text to be omitted; the Word markup may be tracked
        # changes, an unsupported drawing, or a parser mapping issue.
        raise ValueError(
            "STRIKETHROUGH_NOT_DETECTED: Bifa este activă, dar DOCX nu are "
            "caractere tăiate recunoscute în sincronurile selectate. "
            "OPRIT înainte de modificarea V3. Verifică formatările Word "
            "sau debifează pentru a include tot textul.")
    return {"texts":texts,"cutIdsOmitted":all_omitted,
            "removedCharacters":removed,"cutsChanged":changed,
            "policy":"OMIT_STRUCK"}
