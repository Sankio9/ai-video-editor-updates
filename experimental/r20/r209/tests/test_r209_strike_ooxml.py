"""R2.0.9 robust DOCX strike detection: actual OOXML style and tracked revisions.

All documents are synthetic. The source report Bubuieci is not uploaded to
this public repository and therefore cannot be certified from CI alone.
"""
from __future__ import annotations
import sys,tempfile,unittest
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/"candidate"/"tools"))
from r207_docx_strike import extract_strike_policy

class TestR209(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.path=Path(self.tmp.name)/"synthetic.docx"
        self.doc=Document()
    def tearDown(self):self.tmp.cleanup()
    def build(self,text,paragraph=0):
        self.doc.save(self.path)
        return extract_strike_policy(self.path,[{
            "ordinal":2,"paragraph_start":paragraph,"text":text
        }],True)

    def test_direct_strikethrough_preserves_italic_question(self):
        p=self.doc.add_paragraph()
        p.add_run("Vreau să mă dezvolt în profesia dată. ").font.strike=True
        p.add_run("Unde ai vrea să muncești? ").italic=True
        p.add_run("Să fiu bucătar în Anglia.")
        result=self.build(p.text)
        self.assertEqual(result["texts"],["Unde ai vrea să muncești? Să fiu bucătar în Anglia."])
        self.assertGreater(result["removedCharacters"],0)
        self.assertEqual(result["cutsChanged"],1)

    def test_inherited_character_style_strikethrough(self):
        base=self.doc.styles.add_style("Editorial Cut Base",WD_STYLE_TYPE.CHARACTER)
        base.font.strike=True
        derived=self.doc.styles.add_style("Editorial Cut Child",WD_STYLE_TYPE.CHARACTER)
        derived.base_style=base
        p=self.doc.add_paragraph()
        r=p.add_run("Propoziție exclusă. ");r.style=derived
        p.add_run("Întrebare validă.")
        result=self.build(p.text)
        self.assertEqual(result["texts"],["Întrebare validă."])

    def test_explicit_false_override_inherited_true_prevents_omission(self):
        style=self.doc.styles.add_style("Cut Style",WD_STYLE_TYPE.CHARACTER)
        style.font.strike=True
        p=self.doc.add_paragraph()
        good=p.add_run("Textul bun. ");good.style=style;good.font.strike=False
        cut=p.add_run("Text tăiat.");cut.style=style
        result=self.build(p.text)
        self.assertEqual(result["texts"],["Textul bun."])

    def test_inherited_paragraph_style(self):
        base=self.doc.styles.add_style("Deleted Paragraph Base",WD_STYLE_TYPE.PARAGRAPH)
        base.font.strike=True
        child=self.doc.styles.add_style("Deleted Paragraph Child",WD_STYLE_TYPE.PARAGRAPH)
        child.base_style=base
        p=self.doc.add_paragraph(style=child)
        p.add_run("Exclus complet.")
        result=self.build(p.text)
        self.assertEqual(result["texts"],[""])
        self.assertEqual(result["cutIdsOmitted"],["CUT-02"])

    def test_double_strike_with_single_false(self):
        p=self.doc.add_paragraph()
        r=p.add_run("Tăiat dublu. ");r.font.strike=False;r.font.double_strike=True
        p.add_run("Bun.")
        result=self.build(p.text)
        self.assertEqual(result["texts"],["Bun."])

    def test_word_deleted_revision_markup(self):
        p=self.doc.add_paragraph()
        p.add_run("Înainte. ")
        deleted=OxmlElement("w:del")
        deleted.set(qn("w:id"),"1")
        run=OxmlElement("w:r")
        text=OxmlElement("w:delText")
        text.set(qn("xml:space"),"preserve")
        text.text="Replica eliminată. "
        run.append(text);deleted.append(run);p._p.append(deleted)
        p.add_run("Rămas.")
        result=self.build("Înainte. Replica eliminată. Rămas.")
        self.assertEqual(result["texts"],["Înainte. Rămas."])

    def test_revision_text_already_excluded_by_parser_stops_if_no_strike_in_cue(self):
        p=self.doc.add_paragraph()
        deleted=OxmlElement("w:del")
        deleted.set(qn("w:id"),"2")
        run=OxmlElement("w:r")
        dt=OxmlElement("w:delText");dt.text="Text vechi. "
        run.append(dt);deleted.append(run);p._p.append(deleted)
        p.add_run("Text valid.")
        # The editorial parser omitted tracked deletion, so no additional
        # struck characters exist in the target cue. Must not claim omitted.
        with self.assertRaisesRegex(ValueError,"STRIKETHROUGH_NOT_DETECTED"):
            self.build("Text valid.")

    def test_checkbox_on_and_no_strike_stops_before_premiere(self):
        p=self.doc.add_paragraph("Nu există text tăiat.")
        with self.assertRaisesRegex(ValueError,"STRIKETHROUGH_NOT_DETECTED"):
            self.build(p.text)

    def test_keep_original_if_checkbox_off(self):
        p=self.doc.add_paragraph("Propoziție valabilă.")
        self.doc.save(self.path)
        result=extract_strike_policy(self.path,[{"ordinal":1,"paragraph_start":0,"text":p.text}],False)
        self.assertEqual(result["texts"],[p.text])
        self.assertEqual(result["policy"],"INCLUDE_STRUCK")

    def test_case_user_report_subtitle_43_invalid_without_replacement(self):
        p=self.doc.add_paragraph("Vreau să mă dezvolt în profesia dată. Unde ai vrea să muncești?")
        with self.assertRaisesRegex(ValueError,"STRIKETHROUGH_NOT_DETECTED"):
            self.build(p.text)
        self.assertTrue(self.path.is_file())

    def test_ambiguous_cue_does_not_guess(self):
        self.doc.add_paragraph("Aceeași frază. Aceeași frază.")
        with self.assertRaisesRegex(ValueError,"STRIKE_FORMAT_AMBIGUOUS"):
            self.build("Aceeași frază.")

    def test_plain_text_other_cue_but_strike_in_selected_range(self):
        self.doc.add_paragraph("Header fără tăiere.")
        p=self.doc.add_paragraph()
        p.add_run("Exclus!").font.strike=True
        p.add_run(" Inclus.")
        result=self.build(p.text,paragraph=1)
        self.assertEqual(result["texts"],["Inclus."])

    def test_header_struck_but_dialogue_not_struck_stop(self):
        a=self.doc.add_paragraph()
        a.add_run("SINCRON 9931").font.strike=True
        p=self.doc.add_paragraph("Dialog clar.")
        with self.assertRaisesRegex(ValueError,"STRIKETHROUGH_NOT_DETECTED"):
            self.build(p.text,paragraph=1)

    def test_docx_never_modified(self):
        import hashlib
        p=self.doc.add_paragraph()
        p.add_run("Tăiat. ").font.strike=True
        p.add_run("Păstrat.")
        self.doc.save(self.path)
        original=hashlib.sha256(self.path.read_bytes()).hexdigest()
        self.build(p.text)
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).hexdigest(),original)

if __name__=="__main__":unittest.main(verbosity=2)
